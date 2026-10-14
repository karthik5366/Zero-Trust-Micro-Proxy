import os
import docx
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

base_docx = r"C:\Users\karth\.gemini\antigravity\brain\7e02376d-9810-4e77-b16b-5b859704ee5e\.user_uploaded\media_1791397436641_3a0e0420.docx"
img_dir = r"C:\Users\karth\Downloads\Projects\Zero-Trust-Micro-Proxy\report-images\final"
output_docx = r"C:\Users\karth\Downloads\Projects\Zero-Trust-Micro-Proxy\ZetaShield_Final_Report.docx"

print("Loading base document...")
doc = docx.Document(base_docx)

# 1. Cover page logo is LEFT TO THE USER as requested. (No modification to cover logo).

# 2. Update Bonafide Certificate & Examination Date
print("Updating Bonafide Certificate...")
for p in doc.paragraphs[20:35]:
    if "<<FILL>>" in p.text or "held on" in p.text:
        p.text = "The Capstone Project Examination is held on September-October 2026"
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if "Saravanan" in p.text:
        p.text = p.text.replace("Dr. D. Saravanan", "Dr. Ajay Kumar Phulre")
        p.text = p.text.replace("Assistant Professor Sr.", "Professor Sr.")

# Table 0: Program Chair & Supervisor signatures
if len(doc.tables) > 0:
    table0 = doc.tables[0]
    cell0 = table0.rows[0].cells[0]
    cell0.text = "PROGRAM CHAIR,\nDr. Ajay Kumar Phulre, Professor Sr.,\nSchool of Computing Science and Engineering,\nVIT Bhopal University."
    cell1 = table0.rows[0].cells[1]
    cell1.text = "PROJECT SUPERVISOR,\nDr. Sajjad Ahmed, Professor Sr.,\nSchool of Computing Science and Engineering,\nVIT Bhopal University."

# 3. Update Acknowledgement
print("Updating Acknowledgement...")
for p in doc.paragraphs[28:40]:
    if "Saravanan" in p.text:
        p.text = p.text.replace("Dr. D. Saravanan", "Dr. Ajay Kumar Phulre")
        p.text = p.text.replace("Assistant Professor Sr.", "Professor Sr.")

# 4. Map and Insert All 16 Figures
figure_map = {
    4:  ("fig_1_1.png", Inches(6.0)),
    5:  ("fig_1_2.png", Inches(6.0)),
    6:  ("fig_2_1.png", Inches(6.0)),
    8:  ("fig_3_1.png", Inches(6.0)),
    14: ("fig_4_1.png", Inches(6.0)),
    15: ("fig_4_2.png", Inches(6.0)),
    18: ("fig_4_3.png", Inches(6.0)),
    20: ("fig_5_1.png", Inches(5.8)),
    21: ("fig_5_8.png", Inches(5.8)),
    22: ("fig_5_4.png", Inches(5.8)),
    24: ("fig_5_6.png", Inches(5.8)),
    26: ("fig_5_2.png", Inches(5.8)),
    27: ("fig_5_3.png", Inches(5.8)),
    28: ("fig_5_7.png", Inches(6.0)),
    30: ("fig_5_5.png", Inches(5.8)),
    34: ("fig_6_1.png", Inches(6.0)),
}

print("Inserting figures into document tables...")
for t_idx, (img_name, width) in figure_map.items():
    if t_idx < len(doc.tables):
        table = doc.tables[t_idx]
        cell = table.rows[0].cells[0]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_path = os.path.join(img_dir, img_name)
        if os.path.exists(img_path):
            run = p.add_run()
            run.add_picture(img_path, width=width)
            print(f"Inserted Table {t_idx} -> {img_name}")
        else:
            print(f"WARNING: Image not found: {img_path}")

# Remove outer borders of figure tables for a clean visual appearance
def remove_table_borders(table):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'none')
        tblBorders.append(border)
    tblPr.append(tblBorders)

for t_idx in figure_map.keys():
    if t_idx < len(doc.tables):
        remove_table_borders(doc.tables[t_idx])

print(f"Saving updated document to {output_docx}...")
doc.save(output_docx)
print("DOCX successfully saved!")
